# -*- coding: utf-8 -*-
"""workbook_style.py — Ink & Gold header/zebra styling for the workbook."""
from openpyxl.styles import Font, PatternFill, Border, Side

PAL = {
    "accent": "8C6D1F",
    "accent_soft": "C9A84C",
    "ink": "1A1A24",
    "surface_a": "F5F2E8",
    "surface_b": "FFFFFF",
    "line": "D8D2C4",
}

header_font = Font(bold=True, color="FAF7F0", size=11)
header_fill = PatternFill("solid", fgColor=PAL["accent"])
_fill_a = PatternFill("solid", fgColor=PAL["surface_a"])
_fill_b = PatternFill("solid", fgColor=PAL["surface_b"])
_thin = Side(style="thin", color=PAL["line"])
thin_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def zebra_fill(i):
    return _fill_a if i % 2 == 0 else _fill_b
