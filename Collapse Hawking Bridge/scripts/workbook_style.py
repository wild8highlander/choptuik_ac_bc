# -*- coding: utf-8 -*-
"""Local styling helpers for scripts/build_workbook.py (v2.0.1).

Self-contained replacements for the out-of-tree style module used at
2.0.0 (audit advisory IVP-03: the release must not import from paths
outside its own tree, and openpyxl must be declared as an optional
dependency).  Same call signatures as the previous helpers.
"""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT_NAME = "Calibri"

INK = "1F2430"
GOLD = "8A6D1D"
GOLD_SOFT = "F4E9C8"
PAPER = "FFFFFF"
ZEBRA = "F7F5EF"

_MARGIN_W = 2.5


def font_title():
    return Font(name=FONT_NAME, size=14, bold=True, color=INK)


def font_body():
    return Font(name=FONT_NAME, size=10, color=INK)


def align_number():
    return Alignment(horizontal="right", vertical="center")


def align_text():
    return Alignment(horizontal="left", vertical="center")


def setup_sheet(ws, title=None, last_col=None):
    """Standard sheet setup: no grid, margin column, title band at B2."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = _MARGIN_W
    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 6
    if title and last_col:
        ws.merge_cells(start_row=2, start_column=2,
                       end_row=2, end_column=last_col)
        c = ws.cell(row=2, column=2, value=title)
        c.font = font_title()
        c.alignment = Alignment(horizontal="left", vertical="center")


def _paint(ws, row_num, col_start, col_end, fill, font, zebra=None):
    for j in range(col_start, col_end + 1):
        c = ws.cell(row=row_num, column=j)
        c.font = font
        bg = zebra if (zebra and j > col_start) else fill
        if bg:
            c.fill = PatternFill("solid", fgColor=bg)
        if c.alignment.horizontal is None:
            c.alignment = (align_number() if isinstance(c.value, (int, float))
                           else align_text())


def style_header_row(ws, row_num, col_start, col_end):
    _paint(ws, row_num, col_start, col_end, GOLD_SOFT,
           Font(name=FONT_NAME, size=10, bold=True, color=INK))


def style_data_row(ws, row_num, col_start, col_end, row_index):
    _paint(ws, row_num, col_start, col_end, PAPER, font_body(),
           zebra=ZEBRA if row_index % 2 else None)


def style_total_row(ws, row_num, col_start, col_end):
    _paint(ws, row_num, col_start, col_end, GOLD_SOFT,
           Font(name=FONT_NAME, size=10, bold=True, color=GOLD))
